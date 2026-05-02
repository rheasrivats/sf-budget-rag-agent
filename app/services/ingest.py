from __future__ import annotations

import csv
import hashlib
import io
import re
from dataclasses import dataclass
from datetime import datetime
from typing import Iterable
from urllib.parse import urljoin, urlparse

import openpyxl
import requests
from bs4 import BeautifulSoup
from langchain_text_splitters import RecursiveCharacterTextSplitter
from pypdf import PdfReader
from sqlalchemy import func, select
from sqlalchemy.orm import Session

from app.config import settings
from app.models import DocumentChunk, SourceDocument
from rag_embeddings_store import create_embeddings

USER_AGENT = "sf-budget-research-agent/1.0"
SUPPORTED_EXTENSIONS = (".pdf", ".xlsx", ".xlsm", ".csv")
MAX_LINKED_ASSETS = 40
MIN_CHARS = 200


@dataclass
class RawDoc:
    url: str
    title: str
    content_type: str
    text: str
    metadata: dict


def _domain_allowed(url: str) -> bool:
    parsed = urlparse(url)
    return parsed.netloc.endswith("sf.gov")


def _normalize_whitespace(text: str) -> str:
    return re.sub(r"\s+", " ", text).strip()


def _hash_text(text: str) -> str:
    return hashlib.sha256(text.encode("utf-8")).hexdigest()


def _get(url: str) -> requests.Response:
    response = requests.get(url, timeout=40, headers={"User-Agent": USER_AGENT})
    response.raise_for_status()
    return response


def _extract_html(url: str) -> tuple[str, str, list[str]]:
    response = _get(url)
    soup = BeautifulSoup(response.text, "html.parser")
    title = (soup.title.string or "").strip() if soup.title else url
    text = _normalize_whitespace(soup.get_text(" ", strip=True))

    links: list[str] = []
    for a in soup.select("a[href]"):
        href = (a.get("href") or "").strip()
        if not href:
            continue
        absolute = urljoin(url, href)
        if not _domain_allowed(absolute):
            continue
        if absolute.lower().endswith(SUPPORTED_EXTENSIONS):
            links.append(absolute)

    # preserve order, dedupe
    deduped: list[str] = []
    seen: set[str] = set()
    for link in links:
        if link not in seen:
            deduped.append(link)
            seen.add(link)

    return title, text, deduped


def _extract_pdf(url: str) -> str:
    response = _get(url)
    reader = PdfReader(io.BytesIO(response.content))
    parts: list[str] = []
    for page in reader.pages:
        parts.append(page.extract_text() or "")
    return _normalize_whitespace("\n".join(parts))


def _extract_xlsx(url: str) -> str:
    response = _get(url)
    if url.lower().endswith(".csv"):
        decoded = response.content.decode("utf-8", errors="ignore")
        reader = csv.reader(decoded.splitlines())
        rows = [" | ".join(cell.strip() for cell in row if cell and cell.strip()) for row in reader]
        return _normalize_whitespace("\n".join(row for row in rows if row))

    wb = openpyxl.load_workbook(io.BytesIO(response.content), data_only=True)
    rows: list[str] = []
    for sheet in wb.worksheets:
        rows.append(f"Sheet: {sheet.title}")
        for row in sheet.iter_rows(values_only=True):
            values = [str(v).strip() for v in row if v is not None and str(v).strip()]
            if values:
                rows.append(" | ".join(values))
    return _normalize_whitespace("\n".join(rows))


def collect_raw_documents(seed_urls: Iterable[str] = settings.seed_urls) -> list[RawDoc]:
    collected: list[RawDoc] = []
    linked_assets: list[str] = []

    for url in seed_urls:
        if not _domain_allowed(url):
            continue
        try:
            _title, _text, links = _extract_html(url)
        except Exception:
            continue
        # Treat HTML seed pages as discovery pages only. They are useful for
        # finding linked official assets, but should not be chunked/retrieved.
        linked_assets.extend(links)

    # Keep first N assets for predictable runtime/cost in v1.
    asset_urls: list[str] = []
    seen: set[str] = set()
    for link in linked_assets:
        if link in seen:
            continue
        seen.add(link)
        asset_urls.append(link)
        if len(asset_urls) >= MAX_LINKED_ASSETS:
            break

    for url in asset_urls:
        lower = url.lower()
        try:
            if lower.endswith(".pdf"):
                text = _extract_pdf(url)
                content_type = "pdf"
            elif lower.endswith((".xlsx", ".xlsm", ".csv")):
                text = _extract_xlsx(url)
                content_type = "xlsx"
            else:
                continue
        except Exception:
            continue

        if len(text) < MIN_CHARS:
            continue

        collected.append(
            RawDoc(
                url=url,
                title=url.split("/")[-1],
                content_type=content_type,
                text=text,
                metadata={"seed": False},
            )
        )

    return collected


def _chunk_text(text: str) -> list[str]:
    splitter = RecursiveCharacterTextSplitter(
        chunk_size=settings.chunk_size,
        chunk_overlap=settings.chunk_overlap,
    )
    return splitter.split_text(text)


def index_text_document(
    db: Session,
    *,
    url: str,
    title: str,
    content_type: str,
    source_kind: str,
    text: str,
    metadata: dict | None = None,
    embeddings=None,
) -> tuple[SourceDocument, bool]:
    """Upsert one document and rebuild chunks when content changes."""
    metadata = metadata or {}
    content_hash = _hash_text(text)
    source = db.scalar(select(SourceDocument).where(SourceDocument.url == url))

    if source is None:
        source = SourceDocument(
            url=url,
            title=title,
            content_type=content_type,
            source_kind=source_kind,
            content_hash=content_hash,
            raw_text=text,
            metadata_json=metadata,
            active=True,
            last_ingested_at=datetime.utcnow(),
        )
        db.add(source)
        db.flush()
        changed = True
    else:
        source.title = title
        source.content_type = content_type
        source.source_kind = source_kind
        source.metadata_json = metadata
        source.active = True
        source.last_ingested_at = datetime.utcnow()

        changed = source.content_hash != content_hash
        if changed:
            source.content_hash = content_hash
            source.raw_text = text

    if not changed:
        return source, False

    db.query(DocumentChunk).filter(DocumentChunk.source_id == source.id).delete()
    chunks = _chunk_text(text)
    if not chunks:
        return source, True

    embeddings = embeddings or create_embeddings(model=settings.embedding_model)
    vectors = embeddings.embed_documents(chunks)
    for idx, (chunk, vector) in enumerate(zip(chunks, vectors)):
        db.add(
            DocumentChunk(
                source_id=source.id,
                chunk_index=idx,
                content=chunk,
                embedding_json=vector,
                metadata_json={
                    "url": source.url,
                    "title": source.title,
                    "content_type": source.content_type,
                    "source_kind": source.source_kind,
                    **metadata,
                },
            )
        )

    return source, True


def run_ingestion(db: Session) -> dict:
    raw_docs = collect_raw_documents(settings.seed_urls)
    embeddings = create_embeddings(model=settings.embedding_model) if raw_docs else None

    updated_sources = 0

    html_sources = list(
        db.scalars(
            select(SourceDocument).where(
                SourceDocument.content_type == "html",
                SourceDocument.source_kind == "official",
                SourceDocument.active.is_(True),
            )
        )
    )
    for source in html_sources:
        deleted = db.query(DocumentChunk).filter(DocumentChunk.source_id == source.id).delete()
        source.active = False
        source.last_ingested_at = datetime.utcnow()
        if deleted:
            updated_sources += 1

    for raw in raw_docs:
        _source, changed = index_text_document(
            db,
            url=raw.url,
            title=raw.title,
            content_type=raw.content_type,
            source_kind="official",
            text=raw.text,
            metadata=raw.metadata,
            embeddings=embeddings,
        )
        if changed:
            updated_sources += 1

    db.flush()

    total_sources = db.scalar(select(func.count()).select_from(SourceDocument)) or 0
    total_chunks = db.scalar(select(func.count()).select_from(DocumentChunk)) or 0

    return {
        "run_at": datetime.utcnow(),
        "total_sources": int(total_sources),
        "total_chunks": int(total_chunks),
        "updated_sources": updated_sources,
    }
